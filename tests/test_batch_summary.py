import argparse
import asyncio
from unittest.mock import AsyncMock, MagicMock

import pytest
from openpyxl import load_workbook

from pdd_listing_automation import app
from pdd_listing_automation.config import Settings
from pdd_listing_automation.models import Platform, ShopNewArrival
from pdd_listing_automation.pages.upload_platform import AlreadyListedError, UploadRejectedError


@pytest.mark.parametrize("report_fails", [False, True])
def test_mixed_batch_summarizes_results_and_unprocessed_items(tmp_path, monkeypatch, capsys, report_fails):
    settings = Settings(
        root_dir=tmp_path, data_dir=tmp_path / "data", profile_dir=tmp_path / "profile",
        artifacts_dir=tmp_path / "artifacts", database_path=tmp_path / "data/db.sqlite3",
    )
    arrivals = [ShopNewArrival(item_id=str(i), item_url=f"https://market.example.com/item/{i}") for i in range(1, 7)]
    args = argparse.Namespace(
        target_shop="测试店", confirm_publish="PUBLISH", item_ids="1,2,3,4,5,6",
        limit=0, delay_seconds=0, hold_open_seconds=0,
    )

    def result(status, detail=""):
        return ({"title": "商品", "sku": "A", "price": "20", "screenshot_path": "test.png",
                 "listing_status": status, "submission_detail": detail}, None, True)

    upload = AsyncMock(side_effect=[
        result("submitted"), UploadRejectedError("授权已过期"), result("uncertain", "页面关闭"),
        AlreadyListedError("平台重复铺货"), RuntimeError("TargetClosed"),
    ])
    monkeypatch.setattr(app, "run_single_upload", upload)
    monkeypatch.setattr(app, "_record_failure_evidence", AsyncMock())
    monkeypatch.setattr(app, "_close_leaked_upload_pages", AsyncMock())
    if report_fails:
        def fail_export(*args):
            raise PermissionError("文件被占用")
        monkeypatch.setattr(app, "export_workbook_xlsx", fail_export)
    assert asyncio.run(app._process_arrivals(
        MagicMock(), arrivals, args, settings, Platform.PINDUODUO, "test", "商品",
    )) == 1
    assert upload.await_count == 5
    output = capsys.readouterr().out
    assert "本次选择 6 款，已处理 5 款，未处理 1 款" in output
    assert "提交成功 1 款 | 失败 2 款 | 跳过 1 款 | 待核实 1 款" in output
    assert "失败商品 ID：2,5" in output
    assert "待核实商品 ID：3" in output
    assert "提前结束原因：浏览器已关闭" in output
    summary_path = next(settings.artifacts_dir.glob("*-批次总结.txt"))
    assert "授权已过期" in summary_path.read_text(encoding="utf-8")
    if report_fails:
        assert "Excel 保存失败：文件被占用" in output
    else:
        workbook = load_workbook(next(settings.artifacts_dir.glob("*-上架结果.xlsx")))
        try:
            assert workbook.sheetnames == ["本次汇总", "上架结果", "失败清单", "待核实清单", "未处理清单"]
            counts = dict(workbook["本次汇总"].iter_rows(min_row=2, values_only=True))
            assert counts["提交成功"] == 1
            assert counts["失败"] == 2
            assert counts["未处理"] == 1
            assert workbook["未处理清单"]["A2"].value == "6"
        finally:
            workbook.close()

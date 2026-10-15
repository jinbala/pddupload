from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

from pydantic import BaseModel


def export_models(
    models: list[BaseModel],
    output_dir: Path,
    stem: str,
) -> tuple[Path, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    if not models:
        rows: list[dict[str, Any]] = []
    else:
        rows = [model.model_dump(mode="json") for model in models]

    json_path = output_dir / f"{stem}.json"
    csv_path = output_dir / f"{stem}.csv"

    json_path.write_text(
        json.dumps(rows, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    fieldnames = list(rows[0].keys()) if rows else []
    with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        if fieldnames:
            writer.writeheader()
            writer.writerows(rows)

    return json_path, csv_path


def export_workbook_xlsx(
    sheets: list[tuple[str, list[dict[str, Any]]]],
    output_dir: Path,
    stem: str,
) -> Path:
    """把多个 sheet 写成一个真正的 .xlsx（依赖 openpyxl，已列为项目依赖）。

    - 表头加粗；列宽按内容长度粗估，方便直接打开看。
    - sheet 名最长 31 字符（Excel 限制），超出自动截断。
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / f"{stem}.xlsx"

    wb = Workbook()
    wb.remove(wb.active)  # 去掉默认空表，只保留我们创建的

    for sheet_name, rows in sheets:
        ws = wb.create_sheet(title=sheet_name[:31])
        if not rows:
            continue
        fieldnames = list(dict.fromkeys(key for row in rows for key in row))
        ws.append(fieldnames)
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append([row.get(k, "") for k in fieldnames])
        ws.auto_filter.ref = ws.dimensions
        for idx, key in enumerate(fieldnames, 1):
            width = max(len(str(key)), *(len(str(r.get(key, ""))) for r in rows[:500]))
            ws.column_dimensions[get_column_letter(idx)].width = min(max(width + 2, 8), 50)

    wb.save(path)
    return path
